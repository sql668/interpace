"""文档解析服务：把上传的文件转成可喂给模型的文本，并识别扫描件。

按后缀分发：

- `.txt` / `.md`：直接按 UTF-8 读取；
- `.docx`：python-docx 取段落与表格；
- `.pdf`：pypdf 取文字层；文字量过低**且**页面里存在图片对象时判定为「疑似扫描件」。

扫描件**不做 OCR 模型推理**。本项目选的是「渲染成图片交给多模态模型」这条路：
`render_pdf_pages()` 用 pymupdf 把页面渲染成 PNG，由 T3.6 作为多模态内容块送进模型。
之所以不用 MinerU：生产是 2 核 2G 的机器，MinerU 最轻的档位也要独占 2GB 内存，
官方还建议本地部署按 16GB 规划，跑不动。详见 docs/TODO.md 的 T3.0。
"""

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path

import pymupdf
from docx import Document as DocxDocument
from pypdf import PdfReader

# 判定「没有文字层」的阈值：整份文档提取出的正文少于这么多字符，
# 且 PDF 页面里确实存在图片对象，才认定为扫描件。
# 两个条件缺一不可——只看字数会把「损坏的文件」误判成扫描件。
MIN_TEXT_CHARS = 50

# 渲染扫描件时的默认参数。页数与 DPI 都刻意保守：图片 token 很贵，
# 一页扫描件大约几千 token。
DEFAULT_RENDER_DPI = 150
DEFAULT_MAX_RENDER_PAGES = 5

# .docx 表格单元格之间的分隔符，尽量保留表格的可读性。
TABLE_CELL_SEPARATOR = " | "

# 支持解析的后缀。
TEXT_SUFFIXES = {".txt", ".md"}


@dataclass(frozen=True)
class DocumentParseResult:
    """统一的解析结果：三种情况（成功 / 疑似扫描件 / 失败）都能区分。"""

    ok: bool  # 是否拿到了可用文本
    text: str  # 提取出的正文（可能为空）
    suspected_scanned: bool  # 疑似扫描件，需要走多模态图片识别
    message: str  # 给用户看的说明

    @property
    def char_count(self) -> int:
        return len(self.text)


class DocumentParseService:
    """文档解析：按后缀分发，PDF 额外做扫描件识别。"""

    @staticmethod
    def is_supported_document(filename: str) -> bool:
        """是否为支持解析的文档类型（图片 / 音频走别的路径）。"""
        suffix = Path(filename or "").suffix.lower()
        return suffix in TEXT_SUFFIXES or suffix in {".docx", ".pdf"}

    @staticmethod
    def parse(data: bytes, filename: str) -> DocumentParseResult:
        """按后缀解析文件，返回统一结果。

        任何异常都被转成失败结果，不向上抛——解析失败不该影响上传本身，
        文件已经存下来了，只是没读出文本而已。
        """
        try:
            return DocumentParseService._dispatch(data, filename)
        except Exception as exc:
            return DocumentParseResult(
                ok=False,
                text="",
                suspected_scanned=False,
                message=f"解析文件时出错：{exc}",
            )

    @staticmethod
    def _dispatch(data: bytes, filename: str) -> DocumentParseResult:
        suffix = Path(filename or "").suffix.lower()
        if suffix in TEXT_SUFFIXES:
            return DocumentParseService._parse_text(data)
        if suffix == ".docx":
            return DocumentParseService._parse_docx(data)
        if suffix == ".pdf":
            return DocumentParseService._parse_pdf(data)
        return DocumentParseResult(
            ok=False,
            text="",
            suspected_scanned=False,
            message=f"暂不支持解析 {suffix or '该类型'} 文件，请上传 PDF、Word 或纯文本。",
        )

    @staticmethod
    def render_pdf_pages(
        data: bytes,
        max_pages: int = DEFAULT_MAX_RENDER_PAGES,
        dpi: int = DEFAULT_RENDER_DPI,
    ) -> list[bytes]:
        """把 PDF 的前若干页渲染成 PNG 字节，供多模态模型识别（见 T3.6）。"""
        pages: list[bytes] = []
        with pymupdf.open(stream=data, filetype="pdf") as document:
            for index, page in enumerate(document):
                if index >= max_pages:
                    break
                pages.append(page.get_pixmap(dpi=dpi).tobytes("png"))
        return pages

    @staticmethod
    def _parse_text(data: bytes) -> DocumentParseResult:
        text = data.decode("utf-8", errors="replace").strip()
        if not text:
            return DocumentParseResult(False, "", False, "这个文件是空的。")
        return DocumentParseResult(True, text, False, f"已读取文本，共 {len(text)} 字。")

    @staticmethod
    def _parse_docx(data: bytes) -> DocumentParseResult:
        try:
            document = DocxDocument(BytesIO(data))
        except Exception:
            return DocumentParseResult(False, "", False, "这个 Word 文件打不开，可能已损坏。")

        blocks = [p.text.strip() for p in document.paragraphs if p.text.strip()]
        # 简历常用表格排版，表格内容也要取出来，否则会漏掉大段经历
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    blocks.append(TABLE_CELL_SEPARATOR.join(cells))

        text = "\n".join(blocks)
        if not text:
            return DocumentParseResult(False, "", False, "这个 Word 文件里没有可读取的文字。")
        return DocumentParseResult(True, text, False, f"已读取 Word 文档，共 {len(text)} 字。")

    @staticmethod
    def _parse_pdf(data: bytes) -> DocumentParseResult:
        try:
            reader = PdfReader(BytesIO(data))
            pages = [page.extract_text() or "" for page in reader.pages]
        except Exception:
            return DocumentParseResult(False, "", False, "这个 PDF 打不开，可能已损坏或加密。")

        text = "\n".join(part.strip() for part in pages if part.strip()).strip()
        if len(text) >= MIN_TEXT_CHARS:
            return DocumentParseResult(True, text, False, f"已读取 PDF，共 {len(text)} 字。")

        # 文字层几乎为空：再看页面里有没有图片，有就基本可以确定是扫描件
        if DocumentParseService._pdf_has_images(reader):
            return DocumentParseResult(
                ok=False,
                text="",
                suspected_scanned=True,
                message="这份 PDF 没有文字层，像是扫描件或图片导出的，将改用逐页图片识别。",
            )
        if text:
            # 文字少但没有图片：可能是内容本来就很短的电子文档，照样算成功
            return DocumentParseResult(True, text, False, f"已读取 PDF，共 {len(text)} 字。")
        return DocumentParseResult(False, "", False, "这份 PDF 里没有可读取的文字。")

    @staticmethod
    def _pdf_has_images(reader: PdfReader) -> bool:
        """只看页面资源里有没有图片对象，不解码图片内容，成本很低。"""
        for page in reader.pages:
            try:
                resources = page.get("/Resources")
                if resources is None:
                    continue
                xobjects = resources.get_object().get("/XObject")
                if xobjects is None:
                    continue
                for obj in xobjects.get_object().values():
                    try:
                        if obj.get_object().get("/Subtype") == "/Image":
                            return True
                    except Exception:
                        continue
            except Exception:
                continue
        return False
