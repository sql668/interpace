"""测试用的文档夹具：现场生成 PDF / Word，避免往仓库里塞二进制样本。

放在单独模块里，供多个测试文件共用（pytest 会把 tests/ 加到 sys.path）。
"""

from io import BytesIO

import pymupdf
from docx import Document as DocxDocument


def make_text_pdf(text: str) -> bytes:
    """生成一份带文字层的 PDF。"""
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 120), text, fontsize=14)
    data = document.tobytes()
    document.close()
    return data


def make_scanned_pdf(text: str) -> bytes:
    """生成一份「扫描件」：只有位图、没有文字层。

    做法是先把文字页渲染成位图，再把位图插进一张新 PDF——这样文件里
    **有图片对象、没有文字层**，正好对应扫描件的特征。
    """
    source = pymupdf.open()
    page = source.new_page()
    page.insert_text((72, 120), text, fontsize=14)
    image_bytes = page.get_pixmap(dpi=150).tobytes("png")
    source.close()

    scanned = pymupdf.open()
    scanned_page = scanned.new_page()
    scanned_page.insert_image(pymupdf.Rect(0, 0, 595, 842), stream=image_bytes)
    data = scanned.tobytes()
    scanned.close()
    return data


def make_docx(paragraph: str = "张三 · Python 后端工程师") -> bytes:
    """生成一份带表格的 Word 文档（简历常用表格排版）。"""
    document = DocxDocument()
    document.add_paragraph(paragraph)
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "技能"
    table.rows[0].cells[1].text = "FastAPI、MySQL"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
